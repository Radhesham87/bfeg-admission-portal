"""Builds the .xlsx workbook that admins download."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .fields import DOCUMENTS, field_labels

HEADER_FILL = PatternFill("solid", fgColor="12315C")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name="Calibri", size=10)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _columns():
    labels = field_labels()
    cols = ["Reference", "Submitted on"] + list(labels.values())
    for doc in DOCUMENTS:
        cols.append(f"Doc: {doc}")
    return labels, cols


def _row(sub, labels):
    data = sub.data or {}
    docs = sub.documents or {}
    row = [sub.reference,
           sub.created_at.strftime("%d-%m-%Y %H:%M") if sub.created_at else ""]
    row += [data.get(key, "") for key in labels.keys()]
    row += [docs.get(doc, "") for doc in DOCUMENTS]
    return row


def build_workbook(submissions) -> bytes:
    labels, cols = _columns()

    wb = Workbook()
    ws = wb.active
    ws.title = "Applications"
    ws.append(cols)

    for sub in submissions:
        ws.append(_row(sub, labels))

    for c in range(1, len(cols) + 1):
        cell = ws.cell(1, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = min(max(len(cols[c - 1]) + 2, 12), 40)

    for r in range(2, ws.max_row + 1):
        for c in range(1, len(cols) + 1):
            cell = ws.cell(r, c)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="left", vertical="center")

    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}1"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
